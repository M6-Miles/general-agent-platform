import base64
import csv
import hashlib
import io
import json
import logging
import re
import secrets
import time
from collections import OrderedDict
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from uuid import uuid4

from docx import Document as DocxDocument
from fastapi import (
    Depends,
    FastAPI,
    File,
    HTTPException,
    Query,
    Request,
    Response,
    UploadFile,
    status,
)
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

try:  # pypdf is the maintained parser; fallback keeps legacy local environments usable.
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - compatibility for pre-migration environments
    from PyPDF2 import PdfReader
from sqlalchemy import func, text, update
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session
from starlette.exceptions import HTTPException as StarletteHTTPException

from app.audit import record_audit
from app.config import validate_production_secrets
from app.config_secrets import get_settings_with_secrets
from app.db import Base, SessionLocal, engine, get_db
from app.dependencies import Principal, get_principal, require_permission
from app.embeddings import cosine_similarity, get_embedding_provider
from app.errors import (
    http_exception_handler,
    unhandled_exception_handler,
    validation_exception_handler,
)
from app.events import record_runtime_event
from app.idempotency import persist_response, replay_or_reject, request_hash
from app.middleware.rate_limit import RateLimiter
from app.models import (
    Agent,
    AgentVersion,
    Approval,
    AuditEvent,
    Checkpoint,
    KnowledgeBase,
    KnowledgeChunk,
    KnowledgeDocument,
    Run,
    RuntimeEvent,
    RuntimeSnapshot,
    SkillDefinition,
    SystemSetting,
    Tenant,
    TenantQuota,
    ToolCall,
    ToolDefinition,
    ToolVersion,
    VectorMemory,
    WorkflowInstance,
)
from app.observability import (
    observe_document_import,
    observe_embedding_cache_hit,
    observe_rag_query,
    observe_request,
    observe_websocket_connection,
    prometheus_lines,
    timer,
)
from app.providers import get_model_provider
from app.quota import tenant_usage
from app.retrieval import BM25Retriever, HybridRetriever, RetrievalResult
from app.routers.health import router as health_router
from app.runtime import (
    begin_new_attempt,
    expire_pending_approvals,
    transition_run,
    validate_json_schema,
)
from app.schemas import (
    ApprovalDecision,
    ApprovalOut,
    AuditEventOut,
    KnowledgeBaseCreate,
    KnowledgeBaseOut,
    KnowledgeDocumentCreate,
    KnowledgeQuery,
    LoginRequest,
    MemoryCreate,
    MemoryOut,
    RestoreRequest,
    RunCreate,
    RunOut,
    SkillCreate,
    SkillOut,
    SkillReview,
    ToolCallCreate,
    ToolCallOut,
    ToolCreate,
    ToolOut,
    ToolRollback,
    ToolUpdate,
    ToolVersionOut,
    WorkflowInstanceCreate,
    WorkflowInstanceOut,
)
from app.security_utils import redact_json
from app.skill_sdk import (
    content_digest,
    validate_capabilities,
    validate_manifest,
    validate_tool_schemas,
)

try:
    from redis import Redis
except ImportError:  # pragma: no cover
    Redis = None

DEAD_LETTER_STREAM = "agent-runtime:runs:dead-letter"

settings = get_settings_with_secrets()
_PERSISTED_SETTING_KEYS = ("model_name", "embedding_provider", "embedding_model", "rate_limit_enabled", "rate_limit_default")
_LOGIN_LIMIT_WINDOW_SECONDS = 15 * 60
_LOGIN_LIMITS = {"ip": 30, "account": 10, "combo": 5}
_LOGIN_LIMITERS = {
    name: RateLimiter(limit=limit, window_seconds=_LOGIN_LIMIT_WINDOW_SECONDS, redis_url=settings.redis_url)
    for name, limit in _LOGIN_LIMITS.items()
}


def _set_refresh_cookie(response: Response, refresh_token: str) -> None:
    response.set_cookie(
        "agent_refresh_token", refresh_token, httponly=True,
        secure=settings.app_env.lower() in {"production", "prod"}, samesite="strict",
        max_age=7 * 24 * 60 * 60, path="/api/v1/auth",
    )


def _login_limit_keys(request: Request, payload: LoginRequest) -> dict[str, str]:
    ip = request.client.host if request.client else "unknown"
    account = str(payload.email).strip().lower()
    # Keep addresses out of audit metadata and distributed limiter keys.
    ip_key = hashlib.sha256(ip.encode()).hexdigest()
    account_key = hashlib.sha256(account.encode()).hexdigest()
    return {"ip": ip_key, "account": account_key, "combo": f"{ip_key}:{account_key}"}


def _login_rate_limit(request: Request, payload: LoginRequest, *, consume: bool) -> tuple[bool, int]:
    keys = _login_limit_keys(request, payload)
    results = [
        _LOGIN_LIMITERS[name].check(key, limit=_LOGIN_LIMITS[name], consume=consume)
        for name, key in keys.items()
    ]
    blocked = next((result for result in results if not result[0]), None)
    if blocked is not None:
        return False, max(1, blocked[2])
    return True, min((result[2] for result in results), default=_LOGIN_LIMIT_WINDOW_SECONDS)


def _record_login_block(db: Session, request: Request, payload: LoginRequest, retry_after: int) -> None:
    tenant_id = None
    if payload.tenant_slug:
        tenant = db.query(Tenant).filter(Tenant.slug == payload.tenant_slug, Tenant.deleted_at.is_(None)).first()
        tenant_id = tenant.id if tenant else None
    if tenant_id is None:
        return
    record_audit(
        db, tenant_id=tenant_id, actor_id=None, action="user.login_rate_limited",
        resource_type="user", resource_id=str(payload.email).lower(),
        request_id=request.state.request_id,
        metadata={"retryAfterSeconds": retry_after}, outcome="denied",
        reason_code="LOGIN_RATE_LIMITED",
    )
    db.commit()


def load_persisted_settings(db: Session) -> None:
    for key in _PERSISTED_SETTING_KEYS:
        row = db.get(SystemSetting, key)
        if row is not None and isinstance(row.value_json, dict) and "value" in row.value_json:
            setattr(settings, key, row.value_json["value"])
logger = logging.getLogger(__name__)
_started_at = time.monotonic()
_QUERY_EMBEDDING_CACHE: OrderedDict[tuple[str, str], list[float]] = OrderedDict()
_QUERY_EMBEDDING_CACHE_SIZE = 512


def _embedding_metadata() -> tuple[object, str, datetime]:
    provider = get_embedding_provider(settings)
    return provider, getattr(provider, "model_name", "mock"), datetime.now(UTC)


def _embed_query(query: str) -> list[float]:
    """Embed a query with a bounded process-local LRU cache."""
    provider = get_embedding_provider(settings)
    key = (getattr(provider, "model_name", settings.embedding_model), query)
    cached = _QUERY_EMBEDDING_CACHE.get(key)
    if cached is not None:
        _QUERY_EMBEDDING_CACHE.move_to_end(key)
        observe_embedding_cache_hit()
        return cached
    vector = provider.embed(query)
    _QUERY_EMBEDDING_CACHE[key] = vector
    _QUERY_EMBEDDING_CACHE.move_to_end(key)
    if len(_QUERY_EMBEDDING_CACHE) > _QUERY_EMBEDDING_CACHE_SIZE:
        _QUERY_EMBEDDING_CACHE.popitem(last=False)
    return vector


def _retrieve_chunks(chunks: list[KnowledgeChunk], query: str, top_k: int, mode: str, enable_reranker: bool = False) -> list[RetrievalResult]:
    if mode == "bm25_only":
        return BM25Retriever(chunks).search(query, top_k)
    if mode == "hybrid":
        return HybridRetriever(chunks, get_embedding_provider(), enable_reranker=enable_reranker).search(query, top_k)
    query_vector = _embed_query(query)
    ranked = sorted(
        ((cosine_similarity(query_vector, chunk.embedding_json), chunk) for chunk in chunks),
        key=lambda pair: pair[0],
        reverse=True,
    )[:top_k]
    return [RetrievalResult(score, chunk) for score, chunk in ranked]


def _retrieval_candidates(db: Session, tenant_id: str, knowledge_base_id: str, query: str, limit: int = 50, mode: str = "hybrid") -> list[KnowledgeChunk]:
    """Use pgvector for bounded candidate retrieval, with a portable fallback."""
    provider = get_embedding_provider(settings)
    if mode != "bm25_only" and db.bind is not None and db.bind.dialect.name == "postgresql" and getattr(provider, "dimensions", 0) == 384:
        try:
            vector = _embed_query(query)
            literal = "[" + ",".join(str(value) for value in vector) + "]"
            rows = db.execute(
                text(
                    "SELECT id FROM knowledge_chunks "
                    "WHERE tenant_id = :tenant_id AND knowledge_base_id = :knowledge_base_id "
                    "AND embedding_vector IS NOT NULL "
                    "ORDER BY embedding_vector <=> CAST(:embedding AS vector) LIMIT :limit"
                ),
                {"tenant_id": tenant_id, "knowledge_base_id": knowledge_base_id, "embedding": literal, "limit": limit},
            ).all()
            ids = [row.id for row in rows]
            if ids:
                items = db.query(KnowledgeChunk).filter(KnowledgeChunk.id.in_(ids)).all()
                by_id = {item.id: item for item in items}
                return [by_id[item_id] for item_id in ids if item_id in by_id]
        except SQLAlchemyError as exc:
            logger.warning("pgvector retrieval unavailable; using fallback: %s", exc)
    return db.query(KnowledgeChunk).filter(
        KnowledgeChunk.tenant_id == tenant_id,
        KnowledgeChunk.knowledge_base_id == knowledge_base_id,
    ).all()


def _embed_chunks_with_progress(provider: object, chunks: list[str], document_name: str) -> list[list[float]]:
    started = timer()
    vectors = provider.embed_batch(chunks)
    elapsed = timer() - started
    logger.info("knowledge document embedding complete filename=%s chunks=%d elapsed=%.3fs", document_name, len(chunks), elapsed)
    observe_document_import(elapsed)
    return vectors


def _store_native_embedding(db: Session, table: str, item_id: str, vector: list[float]) -> None:
    """Populate pgvector when available; JSON remains the portable fallback."""
    # The schema uses a fixed 384-dimensional pgvector column. Mock embeddings
    # are intentionally smaller, so leave those in the portable JSON column.
    if db.bind is None or db.bind.dialect.name != "postgresql" or len(vector) != 384:
        return
    literal = "[" + ",".join(str(value) for value in vector) + "]"
    db.execute(text(f"UPDATE {table} SET embedding_vector = CAST(:embedding AS vector) WHERE id = :id"), {"embedding": literal, "id": item_id})
_request_count = 0


def _encode_cursor(value: str, item_id: str) -> str:
    raw = json.dumps({"value": value, "id": item_id}, separators=(",", ":")).encode()
    return base64.urlsafe_b64encode(raw).decode().rstrip("=")


def _decode_cursor(cursor: str | None) -> tuple[str, str] | None:
    if not cursor:
        return None
    try:
        padded = cursor + "=" * (-len(cursor) % 4)
        decoded = json.loads(base64.urlsafe_b64decode(padded).decode())
        if not isinstance(decoded, dict) or not isinstance(decoded.get("value"), str) or not isinstance(decoded.get("id"), str):
            raise TypeError
        datetime.fromisoformat(decoded["value"])
        return decoded["value"], decoded["id"]
    except (ValueError, TypeError, OverflowError, json.JSONDecodeError, UnicodeDecodeError):
        raise HTTPException(status_code=400, detail="VALIDATION_ERROR") from None


_RUN_TERMINAL_STATUSES = {"completed", "failed", "cancelled", "timed_out", "budget_exceeded"}


def _runtime_event_stream(run_id: str, tenant_id: str, sequence: int = 0):
    observe_websocket_connection(1)
    try:
        while True:
            with SessionLocal() as stream_db:
                if stream_db.bind is not None and stream_db.bind.dialect.name == "postgresql":
                    stream_db.execute(
                        text("SELECT set_config('app.tenant_id', :tenant_id, true)"),
                        {"tenant_id": tenant_id},
                    )
                events = stream_db.query(RuntimeEvent).filter(
                    RuntimeEvent.run_id == run_id,
                    RuntimeEvent.tenant_id == tenant_id,
                    RuntimeEvent.sequence > sequence,
                ).order_by(RuntimeEvent.sequence).all()
                current_status = stream_db.query(Run.status).filter(
                    Run.id == run_id, Run.tenant_id == tenant_id,
                ).scalar()
                payloads = [
                    (event.event_id, event.event_type, event.sequence, event.payload_json)
                    for event in events
                ]
            for event_id, event_type, event_sequence, event_payload in payloads:
                sequence = event_sequence
                payload = {"data": event_payload, "sequence": event_sequence, "event_id": event_id}
                yield f"id: {event_id}\nevent: {event_type}\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
            if current_status is None or current_status in _RUN_TERMINAL_STATUSES:
                break
            time.sleep(1)
    finally:
        observe_websocket_connection(-1)


def _get_tenant_run_or_404(
    db: Session,
    run_id: str,
    principal: Principal,
    request: Request,
    requested_action: str,
) -> Run:
    run = db.get(Run, run_id)
    if run is None:
        record_audit(
            db,
            tenant_id=principal.tenant_id,
            actor_id=principal.user_id,
            action="permission.denied",
            resource_type="run",
            resource_id=run_id,
            request_id=request.state.request_id,
            metadata={"reason": "resource_not_visible", "requestedAction": requested_action},
            outcome="denied",
        )
        db.commit()
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if run.tenant_id != principal.tenant_id:
        record_audit(
            db,
            tenant_id=principal.tenant_id,
            actor_id=principal.user_id,
            action="permission.denied",
            resource_type="run",
            resource_id=run_id,
            request_id=request.state.request_id,
            metadata={"reason": "cross_tenant_access", "requestedAction": requested_action},
            outcome="denied",
        )
        db.commit()
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    return run


@asynccontextmanager
async def lifespan(_app: FastAPI):
    validate_production_secrets(settings)
    # Reset process-local limiter state on startup/reload so stale test or dev
    # process events cannot affect a fresh application instance.
    from app.dependencies import _RATE_LIMITER
    _RATE_LIMITER.reset()
    for limiter in _LOGIN_LIMITERS.values():
        limiter.reset()
    if settings.app_env == "development" and settings.database_url.startswith("sqlite"):
        Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        load_persisted_settings(db)
    yield


app = FastAPI(title="Auditable Agent Runtime API", version="0.1.0", lifespan=lifespan)
app.add_exception_handler(StarletteHTTPException, http_exception_handler)
app.add_exception_handler(RequestValidationError, validation_exception_handler)
app.add_exception_handler(Exception, unhandled_exception_handler)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Authorization", "Content-Type", "Idempotency-Key", "If-Match"],
)


@app.middleware("http")
async def request_context(request: Request, call_next):
    global _request_count
    _request_count += 1
    request.state.request_id = request.headers.get("X-Request-Id", f"req_{uuid4().hex}")
    started = timer()
    response = await call_next(request)
    observe_request(request.url.path, timer() - started)
    response.headers["X-Request-Id"] = request.state.request_id
    if hasattr(request.state, "rate_limit"):
        limit, remaining, reset = request.state.rate_limit
        response.headers["X-RateLimit-Limit"] = str(limit)
        response.headers["X-RateLimit-Remaining"] = str(remaining)
        response.headers["X-RateLimit-Reset"] = str(reset)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    if settings.app_env == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


@app.get("/metrics", include_in_schema=False)
def metrics() -> Response:
    uptime = time.monotonic() - _started_at
    body = (
        "# TYPE agent_http_requests_total counter\n"
        f"agent_http_requests_total {_request_count}\n"
        "# TYPE agent_process_uptime_seconds gauge\n"
        f"agent_process_uptime_seconds {uptime:.3f}\n"
    )
    body += "\n".join(prometheus_lines()) + ("\n" if prometheus_lines() else "")
    return Response(content=body, media_type="text/plain; version=0.0.4")


app.include_router(health_router)


@app.post("/api/v1/runs", status_code=status.HTTP_202_ACCEPTED)
def create_run(payload: RunCreate, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    idempotency_key = request.headers.get("Idempotency-Key")
    fingerprint = request_hash(request, payload.model_dump(mode="json"))
    replay = replay_or_reject(db, tenant_id=principal.tenant_id, key=idempotency_key, fingerprint=fingerprint)
    if replay is not None:
        return replay
    agent = db.query(Agent).filter(Agent.id == payload.agent_id, Agent.tenant_id == principal.tenant_id, Agent.deleted_at.is_(None)).first()
    if agent is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    published = db.query(AgentVersion).filter(
        AgentVersion.agent_id == agent.id,
        AgentVersion.tenant_id == principal.tenant_id,
        AgentVersion.version == agent.version,
    ).first()
    if published is None or agent.status != "published":
        raise HTTPException(status_code=409, detail="AGENT_VERSION_NOT_PUBLISHED")
    quota = db.query(TenantQuota).filter(TenantQuota.tenant_id == principal.tenant_id).with_for_update().first()
    if quota is None:
        quota = TenantQuota(tenant_id=principal.tenant_id)
        db.add(quota)
        db.flush()
    active_statuses = ("accepted", "preparing", "running", "waiting_approval")
    active_runs = db.query(func.count(Run.id)).filter(Run.tenant_id == principal.tenant_id, Run.status.in_(active_statuses)).scalar() or 0
    if quota.max_concurrent_runs > 0 and active_runs >= quota.max_concurrent_runs:
        raise HTTPException(status_code=429, detail="TENANT_CONCURRENCY_QUOTA_EXCEEDED", headers={"Retry-After": "5"})
    usage = tenant_usage(db, principal.tenant_id)
    if quota.monthly_token_limit > 0:
        used_tokens = int(usage["total_tokens"])
        if used_tokens >= quota.monthly_token_limit:
            raise HTTPException(status_code=429, detail="QUOTA_EXCEEDED", headers={"Retry-After": "3600"})
    if quota.monthly_cost_limit_usd > 0:
        used_cost = float(usage["cost_usd"])
        if used_cost >= quota.monthly_cost_limit_usd:
            raise HTTPException(status_code=429, detail="QUOTA_EXCEEDED", headers={"Retry-After": "3600"})
    if quota.max_workflow_runs_per_day > 0:
        day_start = datetime.now(UTC).replace(hour=0, minute=0, second=0, microsecond=0)
        daily_runs = db.query(func.count(Run.id)).filter(Run.tenant_id == principal.tenant_id, Run.created_at >= day_start).scalar() or 0
        if daily_runs >= quota.max_workflow_runs_per_day:
            raise HTTPException(status_code=429, detail="WORKFLOW_RUN_QUOTA_EXCEEDED", headers={"Retry-After": "3600"})
    run_input = dict(payload.input)
    knowledge_base_id = run_input.get("knowledge_base_id")
    knowledge_query = run_input.get("knowledge_query") or run_input.get("prompt")
    if knowledge_base_id and knowledge_query:
        base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
        if base is None:
            raise HTTPException(status_code=404, detail="NOT_FOUND")
        chunks = _retrieval_candidates(db, principal.tenant_id, base.id, str(knowledge_query), 50, "hybrid")
        ranked = _retrieve_chunks(chunks, str(knowledge_query), 5, "hybrid")
        run_input["retrieved_context"] = [
            {
                "score": round(result.score, 6),
                "content": result.item.content,
                "document_id": result.item.document_id,
                "position": result.item.position,
            }
            for result in ranked
        ]
    definition = published.definition or {}
    budget = definition.get("budget", {})
    run_input = redact_json(run_input)
    run = Run(tenant_id=principal.tenant_id, agent_id=agent.id, agent_version_id=published.id, input_json=run_input, created_by=principal.user_id, budget_json=budget)
    db.add(run)
    db.flush()
    snapshot = RuntimeSnapshot(
        tenant_id=principal.tenant_id,
        run_id=run.id,
        agent_id=agent.id,
        agent_version=published.version,
        schema_version=1,
        definition_json=definition,
        tool_policies_json=definition.get("tool_policies", []),
        approval_policy_json=definition.get("approval_policy", {}),
        budget_policy_json=budget,
    )
    db.add(snapshot)
    run.snapshot_id = snapshot.id
    if payload.workflow_id or definition.get("workflow"):
        db.add(WorkflowInstance(tenant_id=principal.tenant_id, workflow_id=payload.workflow_id or agent.id, run_id=run.id, status="accepted", input_json=run_input, created_by=principal.user_id))
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.accepted", resource_type="run", resource_id=run.id, request_id=request.state.request_id)
    record_runtime_event(db, run=run, event_type="run.accepted", payload={"status": run.status})
    response = {"data": RunOut.model_validate(run).model_dump(mode="json"), "request_id": request.state.request_id}
    persist_response(db, tenant_id=principal.tenant_id, key=idempotency_key, fingerprint=fingerprint, response_status=202, response_json=response, resource_type="run", resource_id=run.id)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        replay = replay_or_reject(db, tenant_id=principal.tenant_id, key=idempotency_key, fingerprint=fingerprint)
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    db.refresh(run)
    try:
        from app.worker import publish_run
        publish_run(run.id, run.tenant_id)
    except Exception as exc:  # noqa: BLE001 - queue is optional and must not fail request acceptance
        logger.warning("redis run publish unavailable: %s", exc)
    return response


@app.get("/api/v1/runs")
def list_runs(request: Request, limit: int = Query(default=20, ge=1, le=100), cursor: str | None = None, status_filter: str | None = Query(default=None, alias="status"), agent_id: str | None = None, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    query = db.query(Run).filter(Run.tenant_id == principal.tenant_id)
    if status_filter:
        query = query.filter(Run.status == status_filter)
    if agent_id:
        query = query.filter(Run.agent_id == agent_id)
    decoded = _decode_cursor(cursor)
    if decoded:
        value, item_id = decoded
        timestamp = datetime.fromisoformat(value)
        query = query.filter((Run.created_at < timestamp) | ((Run.created_at == timestamp) & (Run.id < item_id)))
    runs = query.order_by(Run.created_at.desc(), Run.id.desc()).limit(limit + 1).all()
    has_more = len(runs) > limit
    runs = runs[:limit]
    next_cursor = _encode_cursor(runs[-1].created_at.isoformat(), runs[-1].id) if has_more and runs else None
    return {"data": [RunOut.model_validate(run) for run in runs], "request_id": request.state.request_id, "meta": {"next_cursor": next_cursor}}


@app.get("/api/v1/runs/{run_id}")
def get_run(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    run = _get_tenant_run_or_404(db, run_id, principal, request, "read")
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


@app.post("/api/v1/workflow-instances", status_code=status.HTTP_202_ACCEPTED)
def create_workflow_instance(payload: WorkflowInstanceCreate, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    agent = db.query(Agent).filter(Agent.id == payload.workflow_id, Agent.tenant_id == principal.tenant_id, Agent.deleted_at.is_(None)).first()
    if agent is None or not (agent.definition or {}).get("workflow"):
        raise HTTPException(status_code=404, detail="WORKFLOW_NOT_FOUND")
    run_response = create_run(RunCreate(agent_id=agent.id, input=payload.input, workflow_id=agent.id), request, principal, db)
    run_data = run_response["data"]
    instance = db.query(WorkflowInstance).filter(
        WorkflowInstance.run_id == run_data["id"],
        WorkflowInstance.tenant_id == principal.tenant_id,
    ).one()
    return {"data": WorkflowInstanceOut.model_validate(instance), "request_id": request.state.request_id}


@app.get("/api/v1/workflow-instances")
def list_workflow_instances(request: Request, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    items = db.query(WorkflowInstance).filter(WorkflowInstance.tenant_id == principal.tenant_id).order_by(WorkflowInstance.created_at.desc()).limit(100).all()
    return {"data": [WorkflowInstanceOut.model_validate(item) for item in items], "request_id": request.state.request_id}


@app.get("/api/v1/workflow-instances/{instance_id}")
def get_workflow_instance(instance_id: str, request: Request, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    item = db.query(WorkflowInstance).filter(WorkflowInstance.id == instance_id, WorkflowInstance.tenant_id == principal.tenant_id).first()
    if item is None: raise HTTPException(status_code=404, detail="NOT_FOUND")
    if item.run_id:
        run = db.query(Run).filter(Run.id == item.run_id, Run.tenant_id == principal.tenant_id).first()
        if run and item.status != run.status:
            item.status = run.status
            item.output_json = run.output_json
            db.commit(); db.refresh(item)
    return {"data": WorkflowInstanceOut.model_validate(item), "request_id": request.state.request_id}


@app.post("/api/v1/workflow-instances/{instance_id}/cancel")
def cancel_workflow_instance(instance_id: str, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    item = db.query(WorkflowInstance).filter(WorkflowInstance.id == instance_id, WorkflowInstance.tenant_id == principal.tenant_id).first()
    if item is None: raise HTTPException(status_code=404, detail="NOT_FOUND")
    if item.status in {"completed", "failed", "cancelled"}: raise HTTPException(status_code=409, detail="INVALID_STATE_TRANSITION")
    run = db.query(Run).filter(Run.id == item.run_id, Run.tenant_id == principal.tenant_id).first()
    if run is not None and run.status not in {"completed", "failed", "cancelled", "timed_out", "budget_exceeded"}:
        transition_run(db, run, "cancelled")
        run.finished_at = datetime.now(UTC)
        record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.cancelled", resource_type="run", resource_id=run.id, request_id=request.state.request_id)
    item.status = "cancelled"
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="workflow_instance.cancelled", resource_type="workflow_instance", resource_id=item.id, request_id=request.state.request_id)
    db.commit(); db.refresh(item)
    return {"data": WorkflowInstanceOut.model_validate(item), "request_id": request.state.request_id}


@app.get("/api/v1/conversations", response_model=None)
def list_conversations(request: Request, limit: int = Query(default=20, ge=1, le=100), cursor: str | None = None, status_filter: str | None = Query(default=None, alias="status"), agent_id: str | None = None, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    return list_runs(request, limit, cursor, status_filter, agent_id, principal, db)


@app.get("/api/v1/conversations/{run_id}", response_model=None)
def get_conversation(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    return get_run(run_id, request, principal, db)


@app.get("/api/v1/runs/{run_id}/events")
def run_events(run_id: str, request: Request, principal: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    _get_tenant_run_or_404(db, run_id, principal, request, "events.read")

    last_event_id = request.headers.get("Last-Event-ID")
    sequence = 0
    if last_event_id:
        marker = db.query(RuntimeEvent).filter(RuntimeEvent.run_id == run_id, RuntimeEvent.event_id == last_event_id).first()
        if marker is not None:
            sequence = marker.sequence
        else:
            try:
                sequence = int(last_event_id)
            except ValueError:
                raise HTTPException(status_code=409, detail="RESUME_WINDOW_EXPIRED") from None
        oldest = db.query(func.min(RuntimeEvent.sequence)).filter(RuntimeEvent.run_id == run_id, RuntimeEvent.expires_at > datetime.now(UTC)).scalar()
        if oldest is not None and sequence < oldest - 1:
            raise HTTPException(status_code=409, detail="RESUME_WINDOW_EXPIRED")

    db.close()
    return StreamingResponse(_runtime_event_stream(run_id, principal.tenant_id, sequence), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/api/v1/runs/{run_id}/cancel", status_code=status.HTTP_202_ACCEPTED)
def cancel_run(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    run = db.query(Run).filter(
        Run.id == run_id, Run.tenant_id == principal.tenant_id,
    ).with_for_update().first()
    if run is None:
        _get_tenant_run_or_404(db, run_id, principal, request, "cancel")
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if run.status in {"completed", "failed", "cancelled", "timed_out", "budget_exceeded"}:
        return Response(
            content=json.dumps({"data": RunOut.model_validate(run).model_dump(mode="json"), "request_id": request.state.request_id}),
            status_code=status.HTTP_200_OK,
            media_type="application/json",
        )
    transition_run(db, run, "cancelled")
    run.finished_at = datetime.now(UTC)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.cancelled", resource_type="run", resource_id=run.id, request_id=request.state.request_id)
    db.commit()
    db.refresh(run)
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


def _snapshot_tool(tool: ToolDefinition, actor_id: str) -> ToolVersion:
    return ToolVersion(
        tenant_id=tool.tenant_id,
        tool_definition_id=tool.id,
        version=tool.version,
        name=tool.name,
        description=tool.description,
        executor=tool.executor,
        input_schema=tool.input_schema,
        output_schema=tool.output_schema,
        side_effects=tool.side_effects,
        risk_level=tool.risk_level,
        timeout_ms=tool.timeout_ms,
        manifest_json=tool.manifest_json,
        capabilities_json=tool.capabilities_json,
        content_digest=tool.content_digest,
        signature_status=tool.signature_status,
        created_by=actor_id,
    )


@app.post("/api/v1/tools", status_code=status.HTTP_201_CREATED)
def create_tool(payload: ToolCreate, request: Request, principal: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    values = payload.model_dump()
    try:
        raw_manifest = values.pop("manifest", {})
        manifest = validate_manifest(raw_manifest) if raw_manifest else {}
        capabilities = validate_capabilities(values.pop("capabilities"))
        validate_tool_schemas(values["input_schema"], values["output_schema"])
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    tool = ToolDefinition(
        tenant_id=principal.tenant_id,
        created_by=principal.user_id,
        manifest_json=manifest,
        capabilities_json=capabilities,
        content_digest=content_digest(manifest, values["input_schema"], values["output_schema"]) if manifest else None,
        **values,
    )
    db.add(tool)
    db.flush()  # Get tool.id without committing
    db.add(_snapshot_tool(tool, principal.user_id))
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="tool.created", resource_type="tool", resource_id=tool.id, request_id=request.state.request_id)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    db.refresh(tool)
    return {"data": ToolOut.model_validate(tool), "request_id": request.state.request_id}


@app.get("/api/v1/tools")
def list_tools(request: Request, limit: int = Query(default=100, ge=1, le=100), cursor: str | None = None, principal: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    query = db.query(ToolDefinition).filter(ToolDefinition.tenant_id == principal.tenant_id, ToolDefinition.status == "active")
    decoded = _decode_cursor(cursor)
    if decoded:
        value, item_id = decoded
        query = query.filter((ToolDefinition.name > value) | ((ToolDefinition.name == value) & (ToolDefinition.id > item_id)))
    tools = query.order_by(ToolDefinition.name, ToolDefinition.id).limit(limit + 1).all()
    has_more = len(tools) > limit
    tools = tools[:limit]
    next_cursor = _encode_cursor(tools[-1].name, tools[-1].id) if has_more and tools else None
    return {"data": [ToolOut.model_validate(tool) for tool in tools], "request_id": request.state.request_id, "meta": {"next_cursor": next_cursor}}


@app.get("/api/v1/tools/{tool_id}")
def get_tool(tool_id: str, request: Request, principal: Principal = Depends(require_permission("tool:read")), db: Session = Depends(get_db)):
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
        ToolDefinition.status == "active",
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    return {"data": ToolOut.model_validate(tool), "request_id": request.state.request_id}


@app.patch("/api/v1/tools/{tool_id}")
def update_tool(tool_id: str, payload: ToolUpdate, request: Request, principal: Principal = Depends(require_permission("tool:write")), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
        ToolDefinition.status == "active",
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    expected = request.headers.get("If-Match")
    if expected is None or not expected.strip('"').strip():
        raise HTTPException(status_code=428, detail="PRECONDITION_REQUIRED")
    if expected.strip('"') != str(tool.version):
        raise HTTPException(status_code=412, detail="PRECONDITION_FAILED")
    values = payload.model_dump(exclude_unset=True, exclude_none=True)
    input_schema = values.get("input_schema", tool.input_schema)
    output_schema = values.get("output_schema", tool.output_schema)
    try:
        validate_tool_schemas(input_schema, output_schema)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    for key, value in values.items():
        setattr(tool, key, value)
    tool.version += 1
    if tool.manifest_json:
        tool.content_digest = content_digest(tool.manifest_json, tool.input_schema, tool.output_schema)
        tool.signature_status = "unsigned"
    db.add(_snapshot_tool(tool, principal.user_id))
    record_audit(
        db,
        tenant_id=principal.tenant_id,
        actor_id=principal.user_id,
        action="tool.updated",
        resource_type="tool",
        resource_id=tool.id,
        request_id=request.state.request_id,
        metadata={"version": tool.version},
    )
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    db.refresh(tool)
    return {"data": ToolOut.model_validate(tool), "request_id": request.state.request_id}


@app.get("/api/v1/tools/{tool_id}/versions")
def list_tool_versions(tool_id: str, request: Request, principal: Principal = Depends(require_permission("tool:read")), db: Session = Depends(get_db)):
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
        ToolDefinition.status == "active",
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    versions = db.query(ToolVersion).filter(
        ToolVersion.tool_definition_id == tool_id,
        ToolVersion.tenant_id == principal.tenant_id,
    ).order_by(ToolVersion.version.desc()).all()
    return {"data": [ToolVersionOut.model_validate(version) for version in versions], "request_id": request.state.request_id}


@app.post("/api/v1/tools/{tool_id}/rollback", status_code=status.HTTP_201_CREATED)
def rollback_tool(tool_id: str, payload: ToolRollback, request: Request, principal: Principal = Depends(require_permission("tool:write")), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
        ToolDefinition.status == "active",
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    expected = request.headers.get("If-Match")
    if expected is None or not expected.strip('"').strip():
        raise HTTPException(status_code=428, detail="PRECONDITION_REQUIRED")
    if expected.strip('"') != str(tool.version):
        raise HTTPException(status_code=412, detail="PRECONDITION_FAILED")
    target = db.query(ToolVersion).filter(
        ToolVersion.tool_definition_id == tool_id,
        ToolVersion.tenant_id == principal.tenant_id,
        ToolVersion.version == payload.version,
    ).first()
    if target is None:
        raise HTTPException(status_code=404, detail="VERSION_NOT_FOUND")
    restored_version = target.version
    for field in (
        "name", "description", "executor", "input_schema", "output_schema", "side_effects",
        "risk_level", "timeout_ms", "manifest_json", "capabilities_json", "content_digest",
        "signature_status",
    ):
        setattr(tool, field, getattr(target, field))
    tool.version += 1
    db.add(_snapshot_tool(tool, principal.user_id))
    record_audit(
        db,
        tenant_id=principal.tenant_id,
        actor_id=principal.user_id,
        action="tool.rolled_back",
        resource_type="tool",
        resource_id=tool.id,
        request_id=request.state.request_id,
        metadata={"fromVersion": restored_version, "newVersion": tool.version},
    )
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    db.refresh(tool)
    return {"data": ToolOut.model_validate(tool), "request_id": request.state.request_id}


@app.get("/api/v1/tools/{tool_id}/calls")
def list_tool_calls(tool_id: str, request: Request, limit: int = Query(default=50, ge=1, le=100), principal: Principal = Depends(require_permission("tool:read")), db: Session = Depends(get_db)):
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    calls = db.query(ToolCall).filter(
        ToolCall.tool_definition_id == tool_id,
        ToolCall.tenant_id == principal.tenant_id,
    ).order_by(ToolCall.created_at.desc(), ToolCall.id.desc()).limit(limit).all()
    return {"data": [ToolCallOut.model_validate(call) for call in calls], "request_id": request.state.request_id}


@app.delete("/api/v1/tools/{tool_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_tool(tool_id: str, request: Request, principal: Principal = Depends(require_permission("tool:write")), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    tool = db.query(ToolDefinition).filter(
        ToolDefinition.id == tool_id,
        ToolDefinition.tenant_id == principal.tenant_id,
        ToolDefinition.status == "active",
    ).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    tool.status = "archived"
    record_audit(
        db,
        tenant_id=principal.tenant_id,
        actor_id=principal.user_id,
        action="tool.deleted",
        resource_type="tool",
        resource_id=tool.id,
        request_id=request.state.request_id,
        metadata={"softDelete": True, "version": tool.version},
    )
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.post("/api/v1/skills", status_code=status.HTTP_201_CREATED)
def create_skill(payload: SkillCreate, request: Request, principal: Principal = Depends(require_permission("agent:write")), db: Session = Depends(get_db)):
    try:
        manifest = validate_manifest({"id": payload.slug, "version": payload.version, "description": payload.description, "entrypoint": payload.manifest.get("entrypoint", "skill"), "license": payload.manifest.get("license", "UNLICENSED"), "riskLevel": payload.manifest.get("riskLevel", "low"), "sideEffects": payload.manifest.get("sideEffects", False), **payload.manifest})
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    digest = content_digest(manifest, {}, {})
    skill = SkillDefinition(tenant_id=principal.tenant_id, slug=payload.slug, name=payload.name, description=payload.description, version=payload.version, manifest_json=manifest, content_digest=digest, status="draft", created_by=principal.user_id)
    db.add(skill)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    db.refresh(skill)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="skill.created", resource_type="skill", resource_id=skill.id, request_id=request.state.request_id)
    db.commit()
    return {"data": SkillOut.model_validate(skill), "request_id": request.state.request_id}


@app.get("/api/v1/skills")
def list_skills(request: Request, principal: Principal = Depends(require_permission("agent:read")), db: Session = Depends(get_db)):
    skills = db.query(SkillDefinition).filter(SkillDefinition.tenant_id == principal.tenant_id, SkillDefinition.status != "archived").order_by(SkillDefinition.updated_at.desc()).all()
    return {"data": [SkillOut.model_validate(skill) for skill in skills], "request_id": request.state.request_id}


@app.get("/api/v1/skills/{skill_id}")
def get_skill(skill_id: str, request: Request, principal: Principal = Depends(require_permission("agent:read")), db: Session = Depends(get_db)):
    skill = db.query(SkillDefinition).filter(SkillDefinition.id == skill_id, SkillDefinition.tenant_id == principal.tenant_id).first()
    if skill is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    return {"data": SkillOut.model_validate(skill), "request_id": request.state.request_id}


@app.post("/api/v1/skills/{skill_id}/submit", status_code=status.HTTP_202_ACCEPTED)
def submit_skill(skill_id: str, request: Request, principal: Principal = Depends(require_permission("agent:write")), db: Session = Depends(get_db)):
    skill = db.query(SkillDefinition).filter(SkillDefinition.id == skill_id, SkillDefinition.tenant_id == principal.tenant_id).first()
    if skill is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if skill.status not in {"draft", "rejected"}:
        raise HTTPException(status_code=409, detail="SKILL_ALREADY_SUBMITTED")
    skill.status = "pending_review"; skill.review_reason = None
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="skill.submitted", resource_type="skill", resource_id=skill.id, request_id=request.state.request_id)
    db.commit(); db.refresh(skill)
    return {"data": SkillOut.model_validate(skill), "request_id": request.state.request_id}


@app.post("/api/v1/skills/{skill_id}/review")
def review_skill(skill_id: str, payload: SkillReview, request: Request, principal: Principal = Depends(require_permission("agent:write")), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    skill = db.query(SkillDefinition).filter(SkillDefinition.id == skill_id, SkillDefinition.tenant_id == principal.tenant_id).first()
    if skill is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if skill.status != "pending_review":
        raise HTTPException(status_code=409, detail="SKILL_NOT_PENDING")
    skill.status = "published" if payload.decision == "approve" else "rejected"; skill.review_reason = payload.reason; skill.reviewed_by = principal.user_id
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action=f"skill.{skill.status}", resource_type="skill", resource_id=skill.id, request_id=request.state.request_id)
    db.commit(); db.refresh(skill)
    return {"data": SkillOut.model_validate(skill), "request_id": request.state.request_id}


@app.post("/api/v1/runs/{run_id}/tool-calls", status_code=status.HTTP_202_ACCEPTED)
def create_tool_call(run_id: str, payload: ToolCallCreate, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    run = db.query(Run).filter(Run.id == run_id, Run.tenant_id == principal.tenant_id).first()
    if run is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    existing = db.query(ToolCall).filter(ToolCall.tenant_id == principal.tenant_id, ToolCall.idempotency_key == payload.idempotency_key).first()
    if existing is not None:
        if existing.run_id != run_id:
            raise HTTPException(status_code=409, detail="CONFLICT")
        return {"data": ToolCallOut.model_validate(existing), "request_id": request.state.request_id}
    tool = db.query(ToolDefinition).filter(ToolDefinition.tenant_id == principal.tenant_id, ToolDefinition.name == payload.tool_name, ToolDefinition.status == "active").order_by(ToolDefinition.version.desc()).first()
    if tool is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    try:
        validate_json_schema(payload.input, tool.input_schema or {}, "INVALID_TOOL_INPUT")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    initial_status = "pending_approval" if tool.risk_level in {"high", "critical"} else "created"
    call = ToolCall(tenant_id=principal.tenant_id, run_id=run_id, tool_definition_id=tool.id, input_json=redact_json(payload.input), idempotency_key=payload.idempotency_key, status=initial_status)
    db.add(call)
    db.flush()
    raw_token = None
    if initial_status == "pending_approval":
        raw_token = secrets.token_urlsafe(32)
        db.add(Approval(tenant_id=principal.tenant_id, run_id=run_id, tool_call_id=call.id, token_hash=hashlib.sha256(raw_token.encode()).hexdigest(), snapshot_id=run.snapshot_id, expires_at=datetime.now(UTC) + timedelta(minutes=15)))
        db.flush()
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="tool_call.created", resource_type="tool_call", resource_id=call.id, request_id=request.state.request_id, metadata={"tool": tool.name})
    db.commit()
    db.refresh(call)
    response = {"data": ToolCallOut.model_validate(call).model_dump(mode="json"), "request_id": request.state.request_id}
    if raw_token is not None:
        response["data"]["approval_token"] = raw_token
    return response


@app.post("/api/v1/runs/{run_id}/execute", status_code=status.HTTP_202_ACCEPTED)
def execute_run(run_id: str, request: Request, async_mode: bool = Query(default=False, alias="async"), principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    run = db.query(Run).filter(Run.id == run_id, Run.tenant_id == principal.tenant_id).first()
    if run is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if run.status in {"completed", "failed", "cancelled", "timed_out", "budget_exceeded"}:
        return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}

    if async_mode:
        db.commit()
        try:
            from app.worker import publish_run
            publish_run(run.id, run.tenant_id)
        except Exception as exc:  # noqa: BLE001 - database polling remains available
            logger.warning("redis execute publish unavailable: %s", exc)
        db.expire_all()
        refreshed = db.query(Run).filter(Run.id == run_id, Run.tenant_id == principal.tenant_id).first()
        return {"data": RunOut.model_validate(refreshed or run), "request_id": request.state.request_id}

    # Keep the HTTP endpoint as a compatibility trigger, but delegate execution
    # to the same durable Worker path used by queued runs. This prevents API and
    # Worker from applying different workflow, retry, and tool-call semantics.
    db.commit()
    from app.worker import process_once

    process_once(run_id=run_id)
    db.expire_all()
    refreshed = db.query(Run).filter(Run.id == run_id, Run.tenant_id == principal.tenant_id).first()
    if refreshed is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    return {"data": RunOut.model_validate(refreshed), "request_id": request.state.request_id}

@app.post("/api/v1/runs/{run_id}/execute-stream")
def execute_run_stream(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    """Queue a run and expose its durable lifecycle events as an SSE stream."""
    execute_run(run_id, request, async_mode=True, principal=principal, db=db)
    db.close()
    return StreamingResponse(_runtime_event_stream(run_id, principal.tenant_id), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/api/v1/runs/{run_id}/restore")
def restore_run(run_id: str, payload: RestoreRequest, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    run = _get_tenant_run_or_404(db, run_id, principal, request, "restore")
    checkpoint = db.query(Checkpoint).filter(
        Checkpoint.run_id == run_id,
        Checkpoint.tenant_id == principal.tenant_id,
        Checkpoint.version == payload.checkpoint_version,
    ).first()
    if checkpoint is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if run.status not in {"failed", "cancelled", "waiting_approval"}:
        raise HTTPException(status_code=409, detail="RUN_NOT_RESTORABLE")
    try:
        begin_new_attempt(db, run, checkpoint_version=checkpoint.version, event_type="run.restored", payload={"checkpoint_version": checkpoint.version})
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    run.output_json = {"restored_from": checkpoint.version}
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.restored", resource_type="run", resource_id=run.id, request_id=request.state.request_id, metadata={"checkpointVersion": checkpoint.version})
    db.commit()
    db.refresh(run)
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


@app.post("/api/v1/runs/{run_id}/stop")
def stop_run(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    """Request graceful cancellation; the worker stops before its next node."""
    run = _get_tenant_run_or_404(db, run_id, principal, request, "stop")
    if run.status in {"completed", "failed", "cancelled", "timed_out", "budget_exceeded"}:
        raise HTTPException(status_code=409, detail="RUN_NOT_STOPPABLE")
    if run.status != "stopping":
        try:
            transition_run(db, run, "stopping", payload={"reason": "STOP_REQUESTED"})
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.stop_requested", resource_type="run", resource_id=run.id, request_id=request.state.request_id, metadata={"reason": "STOP_REQUESTED"})
    db.commit()
    db.refresh(run)
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


@app.post("/api/v1/runs/{run_id}/replay", status_code=status.HTTP_202_ACCEPTED)
def replay_run(run_id: str, request: Request, principal: Principal = Depends(require_permission("run:execute")), db: Session = Depends(get_db)):
    """Create a new execution attempt for a failed or dead-lettered Run."""
    run = _get_tenant_run_or_404(db, run_id, principal, request, "replay")
    if run.status not in {"failed", "cancelled", "timed_out", "budget_exceeded"}:
        raise HTTPException(status_code=409, detail="RUN_NOT_REPLAYABLE")
    try:
        begin_new_attempt(db, run)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="run.replayed", resource_type="run", resource_id=run.id, request_id=request.state.request_id, metadata={"attemptId": run.attempt_id})
    db.commit()
    db.refresh(run)
    try:
        from app.worker import publish_run
        publish_run(run.id, run.tenant_id)
    except Exception as exc:  # noqa: BLE001 - database polling remains available
        logger.warning("redis replay publish unavailable: %s", exc)
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


@app.get("/api/v1/approvals")
def list_approvals(request: Request, limit: int = Query(default=100, ge=1, le=100), cursor: str | None = None, principal: Principal = Depends(require_permission("approval:decide")), db: Session = Depends(get_db)):
    expire_pending_approvals(db, tenant_id=principal.tenant_id)
    db.commit()
    query = db.query(Approval).filter(Approval.tenant_id == principal.tenant_id, Approval.status == "pending")
    decoded = _decode_cursor(cursor)
    if decoded:
        value, item_id = decoded
        timestamp = datetime.fromisoformat(value)
        query = query.filter((Approval.created_at > timestamp) | ((Approval.created_at == timestamp) & (Approval.id > item_id)))
    approvals = query.order_by(Approval.created_at, Approval.id).limit(limit + 1).all()
    has_more = len(approvals) > limit
    approvals = approvals[:limit]
    next_cursor = _encode_cursor(approvals[-1].created_at.isoformat(), approvals[-1].id) if has_more and approvals else None
    return {"data": [ApprovalOut.model_validate(item) for item in approvals], "request_id": request.state.request_id, "meta": {"next_cursor": next_cursor}}


@app.get("/api/v1/dlq")
def list_dlq(request: Request, principal: Principal = Depends(get_principal)):
    """List dead-lettered stream entries without exposing another tenant's data."""
    if principal.role not in {"admin", "tenant_admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    if Redis is None:
        return {"data": [], "request_id": request.state.request_id, "meta": {"available": False}}
    try:
        client = Redis.from_url(
            settings.redis_url,
            decode_responses=True,
            socket_connect_timeout=0.5,
            socket_timeout=0.5,
        )
        entries = client.xrevrange(DEAD_LETTER_STREAM, count=100)
    except Exception as exc:
        logger.warning("dead letter stream unavailable: %s", exc)
        raise HTTPException(status_code=503, detail="DLQ_UNAVAILABLE") from exc
    data = [
        {"message_id": message_id, "run_id": values.get("run_id", ""), "source_id": values.get("source_id", ""), "error": values.get("error", "")}
        for message_id, values in entries
        if values.get("tenant_id") == principal.tenant_id
    ]
    return {"data": data, "request_id": request.state.request_id, "meta": {"available": True}}


@app.post("/api/v1/dlq/{message_id}/replay", status_code=status.HTTP_202_ACCEPTED)
def replay_dlq(message_id: str, request: Request, principal: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    if principal.role not in {"admin", "tenant_admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    if Redis is None:
        raise HTTPException(status_code=503, detail="DLQ_UNAVAILABLE")
    try:
        client = Redis.from_url(
            settings.redis_url,
            decode_responses=True,
            socket_connect_timeout=0.5,
            socket_timeout=0.5,
        )
        entries = client.xrange(DEAD_LETTER_STREAM, min=message_id, max=message_id, count=1)
    except Exception as exc:
        raise HTTPException(status_code=503, detail="DLQ_UNAVAILABLE") from exc
    if not entries:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if entries[0][1].get("tenant_id") != principal.tenant_id:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    run_id = entries[0][1].get("run_id", "")
    run = db.query(Run).filter(Run.id == run_id, Run.tenant_id == principal.tenant_id).first()
    if run is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if run.status not in {"failed", "cancelled", "timed_out", "budget_exceeded"}:
        raise HTTPException(status_code=409, detail="RUN_NOT_REPLAYABLE")
    try:
        begin_new_attempt(db, run, event_type="run.replayed", payload={"source": "dlq", "message_id": message_id})
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="dlq.replayed", resource_type="run", resource_id=run.id, request_id=request.state.request_id, metadata={"messageId": message_id, "attemptId": run.attempt_id})
    db.commit()
    try:
        from app.worker import publish_run
        publish_run(run.id, run.tenant_id)
    except Exception as exc:  # noqa: BLE001 - polling fallback remains available
        logger.warning("redis replay publish unavailable: %s", exc)
    return {"data": RunOut.model_validate(run), "request_id": request.state.request_id}


@app.post("/api/v1/approvals/{approval_id}/decision")
def decide_approval(approval_id: str, payload: ApprovalDecision, request: Request, principal: Principal = Depends(require_permission("approval:decide")), db: Session = Depends(get_db)):
    if principal.role not in {"tenant_admin", "admin"}:
        record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="permission.denied", resource_type="approval", resource_id=approval_id, request_id=request.state.request_id, metadata={"requiredRoles": ["admin", "tenant_admin"], "role": principal.role}, outcome="denied")
        db.commit()
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    approval = db.query(Approval).filter(Approval.id == approval_id, Approval.tenant_id == principal.tenant_id).first()
    if approval is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    if approval.status != "pending" or approval.used_at is not None:
        raise HTTPException(status_code=409, detail="APPROVAL_ALREADY_HANDLED")
    expires_at = approval.expires_at.replace(tzinfo=UTC) if approval.expires_at.tzinfo is None else approval.expires_at
    if expires_at <= datetime.now(UTC):
        approval.status = "expired"
        approval.used_at = datetime.now(UTC)
        db.commit()
        raise HTTPException(status_code=409, detail="APPROVAL_ALREADY_HANDLED")
    if approval.token_hash is not None and (payload.token is None or approval.token_hash != hashlib.sha256(payload.token.encode()).hexdigest()):
        raise HTTPException(status_code=409, detail="APPROVAL_ALREADY_HANDLED")
    run = db.get(Run, approval.run_id)
    if run is None or run.snapshot_id != approval.snapshot_id:
        raise HTTPException(status_code=409, detail="VERSION_CONFLICT")
    tool_call = db.get(ToolCall, approval.tool_call_id)
    if tool_call is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    definition = db.get(ToolDefinition, tool_call.tool_definition_id)
    if definition is None or definition.status != "active" or definition.tenant_id != principal.tenant_id:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    next_status = "approved" if payload.decision == "approve" else "rejected"
    consumed_at = datetime.now(UTC)
    consumed = db.execute(
        update(Approval)
        .where(
            Approval.id == approval_id,
            Approval.tenant_id == principal.tenant_id,
            Approval.status == "pending",
            Approval.used_at.is_(None),
            Approval.expires_at > consumed_at,
        )
        .values(status=next_status, decision_reason=payload.reason, decided_by=principal.user_id, used_at=consumed_at)
        .execution_options(synchronize_session=False)
    )
    if consumed.rowcount != 1:
        db.rollback()
        raise HTTPException(status_code=409, detail="APPROVAL_ALREADY_HANDLED")
    db.refresh(approval)
    call = db.get(ToolCall, approval.tool_call_id)
    if call and approval.status == "approved":
        call.status = "created"
        if run is not None and run.status == "waiting_approval":
            transition_run(db, run, "running")
    elif call:
        call.status = "rejected"
        if run is not None and run.status == "waiting_approval":
            transition_run(db, run, "failed")
            run.finished_at = datetime.now(UTC)
            instances = db.query(WorkflowInstance).filter(
                WorkflowInstance.run_id == run.id,
                WorkflowInstance.tenant_id == principal.tenant_id,
            ).all()
            for instance in instances:
                instance.status = "failed"
                instance.output_json = run.output_json
            record_audit(
                db, tenant_id=principal.tenant_id, actor_id=principal.user_id,
                action="run.failed", resource_type="run", resource_id=run.id,
                request_id=request.state.request_id,
                metadata={"reason": "APPROVAL_REJECTED", "approvalId": approval.id},
            )
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action=f"approval.{approval.status}", resource_type="approval", resource_id=approval.id, request_id=request.state.request_id)
    db.commit()
    db.refresh(approval)
    return {"data": ApprovalOut.model_validate(approval), "request_id": request.state.request_id}


@app.get("/api/v1/audit")
def list_audit(request: Request, resource_type: str | None = None, resource_id: str | None = None, limit: int = Query(default=100, ge=1, le=100), cursor: str | None = None, principal: Principal = Depends(require_permission("audit:read")), db: Session = Depends(get_db)):
    query = db.query(AuditEvent).filter(AuditEvent.tenant_id == principal.tenant_id)
    if resource_type:
        query = query.filter(AuditEvent.resource_type == resource_type)
    if resource_id:
        query = query.filter(AuditEvent.resource_id == resource_id)
    decoded = _decode_cursor(cursor)
    if decoded:
        value, item_id = decoded
        timestamp = datetime.fromisoformat(value)
        query = query.filter((AuditEvent.created_at < timestamp) | ((AuditEvent.created_at == timestamp) & (AuditEvent.id < item_id)))
    events = query.order_by(AuditEvent.created_at.desc(), AuditEvent.id.desc()).limit(limit + 1).all()
    has_more = len(events) > limit
    events = events[:limit]
    next_cursor = _encode_cursor(events[-1].created_at.isoformat(), events[-1].id) if has_more and events else None
    return {"data": [AuditEventOut.model_validate(event) for event in events], "request_id": request.state.request_id, "meta": {"next_cursor": next_cursor}}


@app.get("/api/v1/audit/export")
def export_audit(request: Request, principal: Principal = Depends(require_permission("audit:export")), db: Session = Depends(get_db)):
    if principal.role not in {"admin", "tenant_admin"}:
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    events = db.query(AuditEvent).filter(AuditEvent.tenant_id == principal.tenant_id).order_by(AuditEvent.created_at.desc()).limit(5000).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["event_id", "tenant_id", "actor_id", "action", "resource_type", "resource_id", "request_id", "outcome", "created_at"])
    for event in events:
        writer.writerow([event.id, event.tenant_id, event.actor_id or "", event.action, event.resource_type, event.resource_id, event.request_id, event.outcome, event.created_at.isoformat()])
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="audit.exported", resource_type="audit", resource_id=principal.tenant_id, request_id=request.state.request_id, metadata={"count": len(events)})
    db.commit()
    return Response(content=output.getvalue(), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=audit.csv"})


@app.get("/api/v1/costs/summary")
def costs_summary(request: Request, principal: Principal = Depends(require_permission("run:read")), db: Session = Depends(get_db)):
    runs = db.query(Run).filter(Run.tenant_id == principal.tenant_id).all()
    return {"data": {"run_count": len(runs), "total_cost_usd": round(sum(float(run.cost_usd or 0) for run in runs), 8), "total_tokens": sum(int((run.usage_json or {}).get("total_tokens", 0)) for run in runs), "completed": sum(run.status == "completed" for run in runs), "failed": sum(run.status == "failed" for run in runs)}, "request_id": request.state.request_id}


@app.get("/api/v1/tenant/usage")
def tenant_usage_summary(request: Request, principal: Principal = Depends(require_permission("audit:read")), db: Session = Depends(get_db)):
    return {"data": tenant_usage(db, principal.tenant_id), "request_id": request.state.request_id}


@app.post("/api/v1/knowledge-bases", status_code=status.HTTP_201_CREATED)
def create_knowledge_base(payload: KnowledgeBaseCreate, request: Request, principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    item = KnowledgeBase(tenant_id=principal.tenant_id, slug=payload.slug, name=payload.name, description=payload.description, created_by=principal.user_id)
    db.add(item)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="CONFLICT") from exc
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="knowledge_base.created", resource_type="knowledge_base", resource_id=item.id, request_id=request.state.request_id)
    db.commit(); db.refresh(item)
    return {"data": KnowledgeBaseOut.model_validate(item), "request_id": request.state.request_id}


@app.get("/api/v1/knowledge-bases")
def list_knowledge_bases(request: Request, principal: Principal = Depends(require_permission("memory:read")), db: Session = Depends(get_db)):
    items = db.query(KnowledgeBase).filter(KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).order_by(KnowledgeBase.created_at.desc()).all()
    return {"data": [KnowledgeBaseOut.model_validate(item) for item in items], "request_id": request.state.request_id}


@app.get("/api/v1/knowledge-bases/{knowledge_base_id}")
def get_knowledge_base(knowledge_base_id: str, request: Request, principal: Principal = Depends(require_permission("memory:read")), db: Session = Depends(get_db)):
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    documents = db.query(KnowledgeDocument).filter(KnowledgeDocument.knowledge_base_id == base.id, KnowledgeDocument.tenant_id == principal.tenant_id).order_by(KnowledgeDocument.created_at.desc()).all()
    data = KnowledgeBaseOut.model_validate(base).model_dump()
    data["documents"] = []
    for document in documents:
        chunks = db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id, KnowledgeChunk.tenant_id == principal.tenant_id).order_by(KnowledgeChunk.position).all()
        data["documents"].append({
            "id": document.id,
            "filename": document.filename,
            "content_type": document.content_type,
            "status": document.status,
            "created_at": document.created_at,
            "chunks": [{"id": chunk.id, "position": chunk.position, "content": chunk.content, "embedding_dimension": len(chunk.embedding_json or [])} for chunk in chunks],
        })
    data["document_count"] = len(documents)
    return {"data": data, "request_id": request.state.request_id}


@app.delete("/api/v1/knowledge-bases/{knowledge_base_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_knowledge_base(knowledge_base_id: str, request: Request, principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    base.deleted_at = datetime.now(UTC)
    db.query(KnowledgeChunk).filter(KnowledgeChunk.knowledge_base_id == base.id, KnowledgeChunk.tenant_id == principal.tenant_id).delete(synchronize_session=False)
    db.query(KnowledgeDocument).filter(KnowledgeDocument.knowledge_base_id == base.id, KnowledgeDocument.tenant_id == principal.tenant_id).delete(synchronize_session=False)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="knowledge_base.deleted", resource_type="knowledge_base", resource_id=base.id, request_id=request.state.request_id)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.delete("/api/v1/knowledge-bases/{knowledge_base_id}/documents/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_knowledge_document(knowledge_base_id: str, document_id: str, request: Request, principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    document = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id, KnowledgeDocument.knowledge_base_id == knowledge_base_id, KnowledgeDocument.tenant_id == principal.tenant_id).first()
    if document is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id, KnowledgeChunk.tenant_id == principal.tenant_id).delete(synchronize_session=False)
    db.delete(document)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="knowledge_document.deleted", resource_type="knowledge_document", resource_id=document.id, request_id=request.state.request_id)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.post("/api/v1/knowledge-bases/{knowledge_base_id}/documents", status_code=status.HTTP_201_CREATED)
def ingest_knowledge_document(knowledge_base_id: str, payload: KnowledgeDocumentCreate, request: Request, principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None: raise HTTPException(status_code=404, detail="NOT_FOUND")
    quota = db.get(TenantQuota, principal.tenant_id)
    if quota and quota.max_knowledge_documents > 0:
        count = db.query(func.count(KnowledgeDocument.id)).filter(KnowledgeDocument.tenant_id == principal.tenant_id).scalar() or 0
        if count >= quota.max_knowledge_documents:
            raise HTTPException(status_code=429, detail="KNOWLEDGE_DOCUMENT_QUOTA_EXCEEDED", headers={"Retry-After": "3600"})
    document = KnowledgeDocument(tenant_id=principal.tenant_id, knowledge_base_id=base.id, filename=payload.filename, content_type=payload.content_type, created_by=principal.user_id)
    db.add(document); db.flush()
    chunks = [payload.content[i:i + 1200] for i in range(0, len(payload.content), 1200)]
    provider, embedding_model, generated_at = _embedding_metadata()
    vectors = _embed_chunks_with_progress(provider, chunks, payload.filename)
    db.add_all([KnowledgeChunk(tenant_id=principal.tenant_id, knowledge_base_id=base.id, document_id=document.id, position=position, chunk_index=position, content=content, embedding_json=vector, embedding_model=embedding_model, embedding_generated_at=generated_at, metadata_json={"filename": payload.filename, "position": position}) for position, (content, vector) in enumerate(zip(chunks, vectors, strict=True))])
    db.flush()
    for chunk, vector in zip(db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id).all(), vectors, strict=True):
        _store_native_embedding(db, "knowledge_chunks", chunk.id, vector)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="knowledge_document.ingested", resource_type="knowledge_document", resource_id=document.id, request_id=request.state.request_id, metadata={"chunks": len(chunks)})
    db.commit(); db.refresh(document)
    return {"data": {"id": document.id, "filename": document.filename, "status": document.status, "chunks": len(chunks)}, "request_id": request.state.request_id}


@app.get("/api/v1/knowledge-bases/{knowledge_base_id}/search")
def search_knowledge(knowledge_base_id: str, request: Request, q: str = Query(min_length=1, max_length=2000), retrieval_mode: str = Query(default="hybrid", pattern="^(vector_only|bm25_only|hybrid)$"), enable_reranker: bool = False, principal: Principal = Depends(require_permission("memory:read")), db: Session = Depends(get_db)):
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None: raise HTTPException(status_code=404, detail="NOT_FOUND")
    started = timer()
    chunks = _retrieval_candidates(db, principal.tenant_id, base.id, q, 50, retrieval_mode)
    ranked = _retrieve_chunks(chunks, q, 10, retrieval_mode, enable_reranker)
    observe_rag_query(timer() - started)
    return {"data": [{"score": round(result.score, 6), "content": result.item.content, "document_id": result.item.document_id, "position": result.item.position, "chunk_index": result.item.chunk_index} for result in ranked], "request_id": request.state.request_id}


@app.post("/api/v1/knowledge-bases/{knowledge_base_id}/query")
def query_knowledge(knowledge_base_id: str, payload: KnowledgeQuery, request: Request, principal: Principal = Depends(require_permission("memory:read")), db: Session = Depends(get_db)):
    """Return a concise grounded answer together with the retrieved chunks."""
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    started = timer()
    chunks = _retrieval_candidates(db, principal.tenant_id, base.id, payload.query, max(50, payload.top_k), payload.retrieval_mode)
    ranked = _retrieve_chunks(chunks, payload.query, payload.top_k, payload.retrieval_mode, payload.enable_reranker)
    document_ids = {result.item.document_id for result in ranked}
    documents = db.query(KnowledgeDocument).filter(KnowledgeDocument.tenant_id == principal.tenant_id, KnowledgeDocument.id.in_(document_ids)).all() if document_ids else []
    document_names = {document.id: document.filename for document in documents}
    citations = [{"citation_id": index, "document_name": document_names.get(result.item.document_id, "未知文档"), "document_id": result.item.document_id, "chunk_index": result.item.chunk_index, "similarity_score": round(result.score, 6), "content": result.item.content} for index, result in enumerate(ranked, 1)]
    sources = [{"score": citation["similarity_score"], "content": citation["content"], "document_id": citation["document_id"], "position": citation["chunk_index"], "chunk_index": citation["chunk_index"]} for citation in citations]
    if not citations:
        answer = "未找到与问题相关的知识库内容。"
    elif settings.model_provider == "openai_compatible":
        context = "\n\n".join(f"[{citation['citation_id']}] {citation['content']}" for citation in citations)
        prompt = (
            "你是一个严格基于资料回答问题的助手。只能使用下面资料中的事实，"
            "不得编造。回答中的每个事实都必须带有对应的 [数字] 引用；"
            "资料不足时只回答：知识库没有足够依据。\n\n"
            f"资料：\n{context}\n\n问题：{payload.query}"
        )
        try:
            generated = get_model_provider(settings).complete(prompt, {"knowledge_base_id": base.id})
            answer = str(generated.get("text", "")).strip() or "知识库没有足够依据。"
            valid_ids = {f"[{citation['citation_id']}]" for citation in citations}
            cited_ids = set(re.findall(r"\[\d+\]", answer))
            if (cited_ids - valid_ids) or not cited_ids:
                answer = "知识库没有足够依据。"
        except (ValueError, OSError) as exc:
            logger.warning("grounded generation unavailable; using extractive answer: %s", exc)
            answer = "\n\n".join(f"[{citation['citation_id']}] {citation['content']}" for citation in citations[:3])
    else:
        answer = "\n\n".join(f"[{citation['citation_id']}] {citation['content']}" for citation in citations[:3])
    observe_rag_query(timer() - started)
    return {"data": {"query": payload.query, "answer": answer, "sources": sources, "citations": citations}, "request_id": request.state.request_id}


def _extract_document_text(filename: str, content_type: str, raw: bytes) -> str:
    """Extract text from supported uploads while keeping the API bounded and deterministic."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in {"txt", "md", "json"} or content_type.startswith("text/") or content_type == "application/json":
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=415, detail="DOCUMENT_MUST_BE_UTF8_TEXT") from exc
    if suffix == "xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for values in sheet.iter_rows(values_only=True):
                    cells = [str(value).strip() for value in values if value is not None and str(value).strip()]
                    if cells:
                        rows.append(" | ".join(cells))
            workbook.close()
            return "\n".join(rows)
        except Exception as exc:
            raise HTTPException(status_code=422, detail="DOCUMENT_XLSX_PARSE_FAILED") from exc
    if suffix == "pdf" or content_type == "application/pdf":
        try:
            return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
        except Exception as exc:
            raise HTTPException(status_code=422, detail="DOCUMENT_PDF_PARSE_FAILED") from exc
    if suffix in {"docx", "doc"} or content_type in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"}:
        if suffix == "doc":
            raise HTTPException(status_code=415, detail="DOCUMENT_LEGACY_DOC_UNSUPPORTED")
        try:
            return "\n".join(paragraph.text for paragraph in DocxDocument(io.BytesIO(raw)).paragraphs)
        except Exception as exc:
            raise HTTPException(status_code=422, detail="DOCUMENT_DOCX_PARSE_FAILED") from exc
    raise HTTPException(status_code=415, detail="DOCUMENT_FORMAT_UNSUPPORTED")


@app.post("/api/v1/knowledge-bases/{knowledge_base_id}/documents/upload", status_code=status.HTTP_201_CREATED)
async def upload_knowledge_document(knowledge_base_id: str, request: Request, file: UploadFile = File(...), principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    base = db.query(KnowledgeBase).filter(KnowledgeBase.id == knowledge_base_id, KnowledgeBase.tenant_id == principal.tenant_id, KnowledgeBase.deleted_at.is_(None)).first()
    if base is None: raise HTTPException(status_code=404, detail="NOT_FOUND")
    quota = db.get(TenantQuota, principal.tenant_id)
    if quota and quota.max_knowledge_documents > 0:
        count = db.query(func.count(KnowledgeDocument.id)).filter(KnowledgeDocument.tenant_id == principal.tenant_id).scalar() or 0
        if count >= quota.max_knowledge_documents:
            raise HTTPException(status_code=429, detail="KNOWLEDGE_DOCUMENT_QUOTA_EXCEEDED", headers={"Retry-After": "3600"})
    max_upload_bytes = settings.max_upload_bytes
    raw = await file.read(max_upload_bytes + 1)
    if len(raw) > max_upload_bytes:
        raise HTTPException(status_code=413, detail="DOCUMENT_TOO_LARGE")
    content = _extract_document_text(file.filename or "upload.txt", file.content_type or "application/octet-stream", raw)
    if not content.strip(): raise HTTPException(status_code=422, detail="DOCUMENT_EMPTY")
    document = KnowledgeDocument(tenant_id=principal.tenant_id, knowledge_base_id=base.id, filename=file.filename or "upload.txt", content_type=file.content_type or "text/plain", created_by=principal.user_id)
    db.add(document); db.flush(); chunks = [content[i:i + 1200] for i in range(0, len(content), 1200)]
    provider, embedding_model, generated_at = _embedding_metadata()
    vectors = _embed_chunks_with_progress(provider, chunks, document.filename)
    db.add_all([KnowledgeChunk(tenant_id=principal.tenant_id, knowledge_base_id=base.id, document_id=document.id, position=position, chunk_index=position, content=chunk_content, embedding_json=vector, embedding_model=embedding_model, embedding_generated_at=generated_at, metadata_json={"filename": document.filename, "position": position}) for position, (chunk_content, vector) in enumerate(zip(chunks, vectors, strict=True))])
    db.flush()
    for chunk, vector in zip(db.query(KnowledgeChunk).filter(KnowledgeChunk.document_id == document.id).all(), vectors, strict=True):
        _store_native_embedding(db, "knowledge_chunks", chunk.id, vector)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="knowledge_document.uploaded", resource_type="knowledge_document", resource_id=document.id, request_id=request.state.request_id, metadata={"chunks": len(chunks)})
    db.commit()
    return {"data": {"id": document.id, "filename": document.filename, "status": document.status, "chunks": len(chunks)}, "request_id": request.state.request_id}


@app.post("/api/v1/agents/{agent_id}/memories", status_code=status.HTTP_201_CREATED)
def create_memory(agent_id: str, payload: MemoryCreate, request: Request, principal: Principal = Depends(require_permission("memory:write")), db: Session = Depends(get_db)):
    agent = db.query(Agent).filter(Agent.id == agent_id, Agent.tenant_id == principal.tenant_id, Agent.deleted_at.is_(None)).first()
    if agent is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    provider, embedding_model, generated_at = _embedding_metadata()
    embedding = provider.embed(payload.content)
    memory = VectorMemory(tenant_id=principal.tenant_id, agent_id=agent.id, run_id=None, content=payload.content, embedding_json=embedding, embedding_model=embedding_model, embedding_generated_at=generated_at, metadata_json=payload.metadata)
    db.add(memory)
    db.flush()
    _store_native_embedding(db, "vector_memories", memory.id, embedding)
    record_audit(db, tenant_id=principal.tenant_id, actor_id=principal.user_id, action="memory.created", resource_type="memory", resource_id=memory.id, request_id=request.state.request_id)
    db.commit()
    db.refresh(memory)
    return {"data": MemoryOut.model_validate(memory), "request_id": request.state.request_id}


@app.get("/api/v1/agents/{agent_id}/memories/search")
def search_memories(agent_id: str, q: str, request: Request, principal: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    agent = db.query(Agent).filter(Agent.id == agent_id, Agent.tenant_id == principal.tenant_id, Agent.deleted_at.is_(None)).first()
    if agent is None:
        raise HTTPException(status_code=404, detail="NOT_FOUND")
    query_vector = get_embedding_provider().embed(q)
    try:
        vector_literal = "[" + ",".join(str(value) for value in query_vector) + "]"
        rows = db.execute(text("SELECT id, 1 - (embedding_vector <=> CAST(:embedding AS vector)) AS score FROM vector_memories WHERE tenant_id = :tenant_id AND agent_id = :agent_id AND deleted_at IS NULL AND embedding_vector IS NOT NULL ORDER BY embedding_vector <=> CAST(:embedding AS vector) LIMIT 20"), {"embedding": vector_literal, "tenant_id": principal.tenant_id, "agent_id": agent.id}).all()
        items = {item.id: item for item in db.query(VectorMemory).filter(VectorMemory.tenant_id == principal.tenant_id, VectorMemory.agent_id == agent.id, VectorMemory.deleted_at.is_(None), VectorMemory.id.in_([row.id for row in rows])).all()}
        ranked = [(float(row.score), items[row.id]) for row in rows if row.id in items]
    except SQLAlchemyError:
        memories = db.query(VectorMemory).filter(VectorMemory.agent_id == agent.id, VectorMemory.tenant_id == principal.tenant_id, VectorMemory.deleted_at.is_(None)).all()
        ranked = sorted(((cosine_similarity(query_vector, item.embedding_json), item) for item in memories), key=lambda pair: pair[0], reverse=True)[:20]
    return {"data": [{"score": round(score, 6), "memory": MemoryOut.model_validate(item)} for score, item in ranked], "request_id": request.state.request_id}


# Keep route registration at the end so router modules can reuse bootstrap
# helpers without introducing import cycles during application construction.
from app.routers.agents import router as agents_router
from app.routers.auth import router as auth_router
from app.routers.settings import router as settings_router
from app.routers.workflows import router as workflows_router

app.include_router(auth_router)
app.include_router(agents_router)
app.include_router(settings_router)
app.include_router(workflows_router)
