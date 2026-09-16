from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook

try:
    from pypdf import PdfWriter
except ImportError:  # legacy developer environment
    from PyPDF2 import PdfWriter

from app.db import SessionLocal
from app.main import app
from app.models import Tenant, User
from app.security import hash_password


def _text_pdf(text: str) -> bytes:
    """Build a tiny standards-compliant PDF fixture without adding a generator dependency."""
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n",
        b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1", errors="replace")
    objects.append(f"5 0 obj << /Length {len(stream)} >> stream\n".encode() + stream + b"\nendstream endobj\n")
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf)); pdf.extend(obj)
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    pdf.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    pdf.extend(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(pdf)


def login(client: TestClient) -> dict[str, str]:
    response = client.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "ChangeMe123456!"})
    if response.status_code == 401:
        from scripts.seed import main
        main()
        response = client.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "ChangeMe123456!"})
    return {"Authorization": f"Bearer {response.json()['data']['access_token']}"}


def test_knowledge_detail_and_delete_endpoints():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "detail-delete", "name": "详情删除"}).json()["data"]
        document = client.post(f"/api/v1/knowledge-bases/{base['id']}/documents", headers=headers, json={"filename": "guide.txt", "content": "第一段内容"}).json()["data"]
        detail = client.get(f"/api/v1/knowledge-bases/{base['id']}", headers=headers)
        assert detail.status_code == 200
        assert detail.json()["data"]["document_count"] == 1
        assert detail.json()["data"]["documents"][0]["chunks"][0]["embedding_dimension"] > 0
        assert client.delete(f"/api/v1/knowledge-bases/{base['id']}/documents/{document['id']}", headers=headers).status_code == 204
        assert client.get(f"/api/v1/knowledge-bases/{base['id']}", headers=headers).json()["data"]["document_count"] == 0
        assert client.delete(f"/api/v1/knowledge-bases/{base['id']}", headers=headers).status_code == 204
        assert client.get(f"/api/v1/knowledge-bases/{base['id']}", headers=headers).status_code == 404


def test_knowledge_contract_and_multipart_upload_is_tenant_scoped():
    assert "/api/v1/knowledge-bases/{knowledge_base_id}/documents/upload" in app.openapi()["paths"]
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "contract-docs", "name": "合同知识库"})
        assert base.status_code == 201
        base_id = base.json()["data"]["id"]
        uploaded = client.post(f"/api/v1/knowledge-bases/{base_id}/documents/upload", headers=headers, files={"file": ("guide.txt", "租户隔离和审计规则", "text/plain")})
        assert uploaded.status_code == 201
        assert uploaded.json()["data"]["chunks"] == 1
        result = client.get(f"/api/v1/knowledge-bases/{base_id}/search", headers=headers, params={"q": "审计"})
        assert result.status_code == 200
        assert result.json()["data"]


def test_knowledge_query_returns_grounded_answer_and_sources():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "qa-demo", "name": "问答演示"}).json()["data"]
        client.post(f"/api/v1/knowledge-bases/{base['id']}/documents", headers=headers, json={"filename": "rules.txt", "content": "审批流程必须记录审计事件", "content_type": "text/plain"})
        response = client.post(f"/api/v1/knowledge-bases/{base['id']}/query", headers=headers, json={"query": "审批审计", "top_k": 5})
        assert response.status_code == 200
        data = response.json()["data"]
        assert data["answer"]
        assert data["sources"][0]["content"] == "审批流程必须记录审计事件"
        assert data["citations"][0]["citation_id"] == 1
        assert data["citations"][0]["document_name"] == "rules.txt"
        assert "[1]" in data["answer"]


def test_generative_rag_rejects_unsupported_citations(monkeypatch):
    import app.main as main_module

    class FakeProvider:
        def complete(self, prompt, metadata):
            return {"text": "这是编造的结论 [99]"}

    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "rag-invalid-citation", "name": "非法引用"}).json()["data"]
        client.post(f"/api/v1/knowledge-bases/{base['id']}/documents", headers=headers, json={"filename": "rules.txt", "content": "审批必须记录审计事件", "content_type": "text/plain"})
        monkeypatch.setattr(main_module.settings, "model_provider", "openai_compatible")
        monkeypatch.setattr(main_module, "get_model_provider", lambda settings: FakeProvider())
        response = client.post(f"/api/v1/knowledge-bases/{base['id']}/query", headers=headers, json={"query": "审批审计", "top_k": 3})
        assert response.status_code == 200
        assert response.json()["data"]["answer"] == "知识库没有足够依据。"


def test_generative_rag_falls_back_to_extractive_answer_on_provider_failure(monkeypatch):
    import app.main as main_module

    class BrokenProvider:
        def complete(self, prompt, metadata):
            raise OSError("provider unavailable")

    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "rag-provider-failure", "name": "生成失败回退"}).json()["data"]
        client.post(f"/api/v1/knowledge-bases/{base['id']}/documents", headers=headers, json={"filename": "rules.txt", "content": "审批必须记录审计事件", "content_type": "text/plain"})
        monkeypatch.setattr(main_module.settings, "model_provider", "openai_compatible")
        monkeypatch.setattr(main_module, "get_model_provider", lambda settings: BrokenProvider())
        response = client.post(f"/api/v1/knowledge-bases/{base['id']}/query", headers=headers, json={"query": "审批审计", "top_k": 3})
        assert response.status_code == 200
        assert response.json()["data"]["answer"].startswith("[1]")


def test_query_embedding_cache_performance(monkeypatch):
    import app.main as main_module

    main_module._QUERY_EMBEDDING_CACHE.clear()
    calls = 0

    class Provider:
        model_name = "test-cache"

        def embed(self, query):
            nonlocal calls
            calls += 1
            return [1.0, 0.0]

    monkeypatch.setattr(main_module, "get_embedding_provider", lambda settings=None: Provider())
    assert main_module._embed_query("重复问题") == [1.0, 0.0]
    assert main_module._embed_query("重复问题") == [1.0, 0.0]
    assert calls == 1


def test_run_persists_tenant_scoped_retrieved_context():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "run-context", "name": "运行上下文"}).json()["data"]
        client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("context.txt", "审批流程必须记录审计事件", "text/plain")})
        agent = client.post("/api/v1/agents", headers=headers, json={"name": "context-agent", "definition": {}}).json()["data"]
        assert client.post(f"/api/v1/agents/{agent['id']}/publish", headers=headers).status_code == 201
        run = client.post("/api/v1/runs", headers=headers, json={"agent_id": agent["id"], "input": {"prompt": "审计", "knowledge_base_id": base["id"]}})
        assert run.status_code == 202
        detail = client.get(f"/api/v1/runs/{run.json()['data']['id']}", headers=headers).json()["data"]
        assert detail["input_json"]["retrieved_context"][0]["content"]


def test_knowledge_base_isolation_and_readonly_permissions():
    with SessionLocal() as db:
        tenant = Tenant(name="隔离租户", slug="isolated-tenant")
        db.add(tenant)
        db.flush()
        db.add(User(tenant_id=tenant.id, email="isolated@example.com", display_name="隔离管理员", role="tenant_admin", password_hash=hash_password("ChangeMe123456!")))
        db.commit()

    with TestClient(app) as client:
        owner = login(client)
        created = client.post("/api/v1/knowledge-bases", headers=owner, json={"slug": "private", "name": "私有库"})
        assert created.status_code == 201
        base_id = created.json()["data"]["id"]
        assert client.post(f"/api/v1/knowledge-bases/{base_id}/documents/upload", headers=owner, files={"file": ("private.txt", "只属于原租户的内容", "text/plain")}).status_code == 201

        isolated_login = client.post("/api/v1/auth/login", json={"email": "isolated@example.com", "password": "ChangeMe123456!"})
        assert isolated_login.status_code == 200
        isolated = {"Authorization": f"Bearer {isolated_login.json()['data']['access_token']}"}
        assert client.get(f"/api/v1/knowledge-bases/{base_id}/search", headers=isolated, params={"q": "内容"}).status_code == 404
        assert client.post(f"/api/v1/knowledge-bases/{base_id}/documents/upload", headers=isolated, files={"file": ("cross.txt", "越权", "text/plain")}).status_code == 404

        readonly_login = client.post("/api/v1/auth/login", json={"email": "readonly@example.com", "password": "ChangeMe123456!"})
        readonly = {"Authorization": f"Bearer {readonly_login.json()['data']['access_token']}"}
        assert client.get("/api/v1/knowledge-bases", headers=readonly).status_code == 200
        assert client.post("/api/v1/knowledge-bases", headers=readonly, json={"slug": "no-write", "name": "禁止写入"}).status_code == 403


def test_pdf_and_docx_uploads_are_parsed_into_searchable_chunks():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "office-files", "name": "办公文档"}).json()["data"]
        pdf_stream = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.write(pdf_stream)
        pdf = client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("empty.pdf", pdf_stream.getvalue(), "application/pdf")})
        assert pdf.status_code == 422

        doc_stream = BytesIO()
        document = DocxDocument()
        document.add_paragraph("DOCX 审计操作必须保留租户上下文")
        document.save(doc_stream)
        docx = client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("audit.docx", doc_stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert docx.status_code == 201
        assert docx.json()["data"]["chunks"] == 1
        search = client.get(f"/api/v1/knowledge-bases/{base['id']}/search", headers=headers, params={"q": "租户上下文"})
        assert search.status_code == 200
        assert "DOCX" in search.json()["data"][0]["content"]


def test_text_pdf_upload_is_searchable():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "text-pdf", "name": "文本 PDF"}).json()["data"]
        payload = _text_pdf("Tenant isolation PDF audit rules")
        uploaded = client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("rules.pdf", payload, "application/pdf")})
        assert uploaded.status_code == 201, uploaded.text
        result = client.get(f"/api/v1/knowledge-bases/{base['id']}/search", headers=headers, params={"q": "audit"})
        assert result.status_code == 200
        assert "PDF" in result.json()["data"][0]["content"]


def test_xlsx_upload_is_parsed_into_searchable_chunks():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "xlsx-docs", "name": "XLSX 文档"}).json()["data"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rules"
        sheet.append(["tenant isolation", "must be audited"])
        stream = BytesIO()
        workbook.save(stream)
        uploaded = client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("rules.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert uploaded.status_code == 201, uploaded.text
        assert uploaded.json()["data"]["chunks"] == 1
        result = client.get(f"/api/v1/knowledge-bases/{base['id']}/search", headers=headers, params={"q": "tenant isolation"})
        assert result.status_code == 200
        assert "tenant isolation" in result.text


def test_hybrid_search_mode_returns_keyword_match():
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "hybrid-search", "name": "Hybrid"}).json()["data"]
        client.post(
            f"/api/v1/knowledge-bases/{base['id']}/documents/upload",
            headers=headers,
            files={"file": ("api.txt", "FastAPI endpoint uses PostgreSQL", "text/plain")},
        )
        result = client.get(
            f"/api/v1/knowledge-bases/{base['id']}/search",
            headers=headers,
            params={"q": "PostgreSQL", "retrieval_mode": "hybrid"},
        )
        assert result.status_code == 200
        assert result.json()["data"][0]["content"] == "FastAPI endpoint uses PostgreSQL"


def test_worker_model_receives_retrieved_context(monkeypatch):
    from app import worker
    observed = {}
    with TestClient(app) as client:
        headers = login(client)
        base = client.post("/api/v1/knowledge-bases", headers=headers, json={"slug": "worker-context", "name": "Worker 上下文"}).json()["data"]
        client.post(f"/api/v1/knowledge-bases/{base['id']}/documents/upload", headers=headers, files={"file": ("context.txt", "检索上下文必须传递给模型", "text/plain")})
        agent = client.post("/api/v1/agents", headers=headers, json={"name": "worker-context-agent", "definition": {"workflow": [{"key": "model", "type": "model"}]}}).json()["data"]
        assert client.post(f"/api/v1/agents/{agent['id']}/publish", headers=headers).status_code == 201
        run_id = client.post("/api/v1/runs", headers=headers, json={"agent_id": agent["id"], "input": {"prompt": "上下文", "knowledge_base_id": base["id"]}}).json()["data"]["id"]
    def fake_model(prompt, context):
        observed.update(context)
        return {"text": "ok"}, {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2, "estimated_cost_usd": 0}
    monkeypatch.setattr(worker, "complete_model", fake_model)
    with SessionLocal() as db:
        persisted = db.get(worker.Run, run_id)
        persisted.status = "running"
        db.commit()
        worker.execute_workflow(db, persisted, db.get(worker.Agent, persisted.agent_id))
    assert observed["retrieved_context"][0]["content"]
